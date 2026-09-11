from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path


OUT = Path(__file__).with_name("RWA_产品介绍_竞品比较与市场潜力评估_合作方版.docx")
FONT = "Hiragino Sans GB"

NAVY = "17324D"
BLUE = "2E74B5"
TEAL = "147D74"
GOLD = "A67525"
INK = "20262E"
MUTED = "5C6670"
LIGHT = "F4F6F9"
BLUEGRAY = "E8EEF5"
PALE_GREEN = "EAF5F2"
PALE_GOLD = "FFF5DE"
RED = "9B1C1C"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, dxa):
    tcPr = cell._tc.get_or_add_tcPr()
    tcW = tcPr.find(qn("w:tcW"))
    if tcW is None:
        tcW = OxmlElement("w:tcW")
        tcPr.append(tcW)
    tcW.set(qn("w:w"), str(dxa))
    tcW.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_fixed(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    tblW.set(qn("w:w"), str(sum(widths)))
    tblW.set(qn("w:type"), "dxa")
    tblLayout = tblPr.find(qn("w:tblLayout"))
    if tblLayout is None:
        tblLayout = OxmlElement("w:tblLayout")
        tblPr.append(tblLayout)
    tblLayout.set(qn("w:type"), "fixed")
    tblInd = tblPr.find(qn("w:tblInd"))
    if tblInd is None:
        tblInd = OxmlElement("w:tblInd")
        tblPr.append(tblInd)
    tblInd.set(qn("w:w"), "120")
    tblInd.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for w in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(w))
        grid.append(gc)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            set_cell_width(cell, widths[i])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    trPr.append(cant_split)


def set_run_font(run, size=11, bold=None, italic=None, color=INK, font=FONT):
    run.font.name = font
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), font)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    new_run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), FONT)
    rpr.append(rfonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "19")
    rpr.append(sz)
    new_run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_para(doc, text="", *, size=11, bold=False, italic=False, color=INK,
             align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=8, line=1.333, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    for r in p.runs:
        set_run_font(r, size=10.6)
    if not p.runs:
        r = p.add_run(text)
        set_run_font(r, size=10.6)
    else:
        p.runs[0].text = text
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.208
    r = p.add_run(text)
    set_run_font(r, size=10.6)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size={1:16, 2:13, 3:12}[level], bold=True,
                 color=BLUE if level < 3 else NAVY)
    return p


def add_callout(doc, title, body, fill=PALE_GREEN, accent=TEAL):
    t = doc.add_table(rows=1, cols=1)
    set_table_fixed(t, [9360])
    cell = t.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_run_font(r, size=11.5, bold=True, color=accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.2
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.4, color=INK)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def add_table(doc, headers, rows, widths, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_row_cant_split(hdr)
    for i, h in enumerate(headers):
        set_cell_shading(hdr.cells[i], NAVY)
        p = hdr.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        set_run_font(r, size=font_size, bold=True, color=WHITE)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, val in enumerate(row):
            if ri % 2 == 1:
                set_cell_shading(cells[i], LIGHT)
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(val))
            set_run_font(r, size=font_size, color=INK, bold=(i == 0))
    set_table_fixed(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def style_document(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = Inches(1)
    sec.bottom_margin = Inches(1)
    sec.left_margin = Inches(1)
    sec.right_margin = Inches(1)
    sec.header_distance = Inches(0.492)
    sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.333
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, NAVY, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(color)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        st = doc.styles[name]
        st.font.name = FONT
        st.font.size = Pt(10.6)
        st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.line_spacing = 1.208


def add_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("RWA TRUST & LIFECYCLE NETWORK  |  合作方材料")
    set_run_font(r, size=8.5, bold=True, color=MUTED)
    pPr = hp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "D7DBE2")
    pBdr.append(bottom)
    pPr.append(pBdr)

    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r1 = fp.add_run("合作讨论稿  ·  2026年8月  |  第 ")
    set_run_font(r1, size=8.5, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    fp._p.append(fld)
    r2 = fp.add_run(" 页")
    set_run_font(r2, size=8.5, color=MUTED)


def add_cover(doc):
    add_para(doc, "", after=44)
    add_para(doc, "合作方产品与市场评估", size=11, bold=True, color=GOLD,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
    add_para(doc, "RWA可信流转与隐私结算平台", size=28, bold=True, color=NAVY,
             align=WD_ALIGN_PARAGRAPH.CENTER, after=10, line=1.0)
    add_para(doc, "产品作用介绍 · 主流竞品比较 · 架构先进性判断 · 市场潜力评估",
             size=14, color=TEAL, align=WD_ALIGN_PARAGRAPH.CENTER, after=34, line=1.15)
    add_para(doc, "面向香港及迪拜的持牌机构、资产管理人、基金行政、分销渠道、托管与技术合作方",
             size=10.5, italic=True, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=56)
    add_callout(doc, "一句话定位",
                "把发行人、分销/KYC、投资者、基金行政/名册、托管、审计监管与隐私结算连接为一套可验证、可回滚、可追责的机构级控制平面；它不是新建一个交易所，也不以代码替代牌照或法律终局。",
                fill=PALE_GOLD, accent=GOLD)
    add_para(doc, "版本：合作方讨论稿 v1.1  |  日期：2026年8月26日",
             size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=36, after=4)
    add_para(doc, "重要边界：当前已达到合成数据机构 PoC 状态；真实资产上线仍依赖持牌主体、法定名册/银行/托管连接、生产 KMS/HSM 与正式 ZK 审计及可信设置。",
             size=9.2, color=RED, align=WD_ALIGN_PARAGRAPH.CENTER, after=0, line=1.2)
    doc.add_page_break()


def add_contents(doc):
    add_heading(doc, "目录与阅读建议", 1)
    rows = [
        ("01", "执行摘要", "先看结论、竞争位置和合作价值"),
        ("02", "行业痛点与产品作用", "理解为什么不是普通代币发行工具"),
        ("03", "产品架构与核心流程", "理解控制面、账本、隐私与法定终局的关系"),
        ("04", "竞品比较与先进性判断", "按维度判断，而非笼统宣称领先"),
        ("05", "市场潜力与区域机会", "全球规模信号、香港/迪拜窗口与可服务市场"),
        ("06", "商业模式与客户切入", "谁付费、为何付费、如何形成首个项目"),
        ("07", "当前成熟度、缺口与路线图", "区分已完成、外部依赖和生产阻断项"),
        ("08", "合作建议与结论", "建议的 PoC 范围和双方投入"),
    ]
    add_table(doc, ["章节", "主题", "阅读目的"], rows, [850, 2750, 5760], font_size=9.6)
    add_callout(doc, "本报告的判断口径",
                "“架构先进”只表示某个设计维度更强，不等于产品整体更成熟。公开资料无法证明竞品未披露的内部实现，因此比较采用“公开能力 + 本地代码证据”的保守口径。市场预测差异极大，报告将预测视为情景，不视为确定收入。")
    doc.add_page_break()


def build():
    doc = Document()
    style_document(doc)
    add_cover(doc)
    add_contents(doc)

    add_heading(doc, "1. 执行摘要", 1)
    add_callout(doc, "核心结论",
                "赛道潜力高、区域政策窗口真实，但竞争并不空白。我们的最佳定位不是“又一个发币平台”，而是服务持牌机构的 RWA 可信流转与隐私结算控制平面：把资格、业务状态、三账一致性、密码学证明、权威根、外部执行回执和法定名册终局分层管理。")
    add_para(doc, "从市场竞争看，Securitize、Tokeny、Taurus、Canton 等产品已经分别占据持牌全栈、合规代币标准、多链发行/托管、隐私互操作网络等高地。我们的架构不能诚实地被称为“整体上比所有主流产品更先进”；但在“隐私转让 + 数据最小披露 + 证明终局/资产终局/法定终局分离 + PostgreSQL 原子账本 + 外部事实双人确认”的组合上，形成了明确且可出售的差异化。")
    add_para(doc, "市场潜力应分三层理解：全球 tokenized RWA 仍处早期，公开预测从 2030 年约 2 万亿美元到 14 万亿美元不等，显示方向一致但口径和不确定性极大；香港和迪拜均已形成监管沙盒、代币化产品规则或真实价值试点；真正可获取的收入取决于能否获得持牌机构、基金行政/名册、银行与托管连接，而不是取决于电路本身。")
    add_table(doc, ["判断项", "结论", "含义"], [
        ("产品价值", "高", "解决多角色协同、终局错配、隐私泄露与审计追责，而非只做链上铸币"),
        ("技术差异化", "中高", "组合架构有辨识度；正式审计、性能、灾备和生产连接尚待完成"),
        ("市场窗口", "高但早期", "政策与机构试点加速，真实 RWA 规模仍远小于预测"),
        ("竞争强度", "高", "头部平台有牌照、客户、链与托管生态；正面替代成本高"),
        ("建议切入", "机构中间层", "与持牌机构合作，提供控制面、隐私结算和一致性证明模块"),
    ], [1600, 1450, 6310], font_size=9.4)

    add_heading(doc, "2. 行业痛点与产品作用", 1)
    add_heading(doc, "2.1 RWA 项目真正难的不是“把资产写到链上”", 2)
    for x in [
        "权利来源分散：发行人、基金行政、过户代理、托管、银行与链上账本各有一套事实，任何一边不同步都可能产生错误所有权。",
        "终局概念混乱：proof 通过、链上交易确认、资产进入权威状态树、银行交收、法定名册更新并不是同一时刻。",
        "隐私与可审计冲突：商业金额、持仓与交易关系不能全网公开，但监管、审计和运营又必须能够证明规则被执行。",
        "合规不是一个布尔值：KYC/AML、投资者类别、司法辖区、销售限制、产品暂停、额度、NAV 时效和例外处置都在生命周期中持续变化。",
        "机构系统以异步消息连接：回调丢失、重放、乱序、重复、超时和人工例外是常态，不能用“链上成功”掩盖链下失败。",
    ]:
        add_bullet(doc, x)
    add_heading(doc, "2.2 我们解决什么", 2)
    add_table(doc, ["角色", "典型痛点", "平台提供的能力"], [
        ("发行人/资产管理人", "产品规则、份额与渠道状态难统一", "产品规则版本化、暂停/恢复、限额、NAV 证据、全生命周期状态机"),
        ("分销/KYC", "资格变化无法及时传导", "服务端身份、MFA、权限、凭证时效与每次交易重新校验"),
        ("投资者", "申购/转让/赎回状态不透明，隐私易泄露", "最小披露回执、受控隐私转让、失败不改变余额"),
        ("基金行政/名册", "链上份额与法定记录可能错位", "独立资产账与名册账、逐资产守恒、权威回调与异常闭环"),
        ("托管/银行", "外部执行与平台状态难原子衔接", "签名回调、顺序控制、Outbox、幂等、对账和双人确认"),
        ("监管/审计", "既要可追责又不能看到无关明文", "追加式审计链、角色化脱敏视图、证据包和不可变最终回执"),
    ], [1450, 3100, 4810], font_size=8.8)

    add_heading(doc, "3. 产品架构与核心流程", 1)
    add_heading(doc, "3.1 四层架构", 2)
    add_table(doc, ["层", "主要组件", "职责边界"], [
        ("业务控制层", "角色、规则、状态机、资格、NAV、例外", "决定交易是否可进入执行；不自称法律意见"),
        ("可信账本层", "PostgreSQL 原子事务、资产账、名册账、审计、Outbox", "保证内部状态、账本、回执与消息同成同败"),
        ("隐私结算层", "Groth16 JoinSplit、Merkle root、nullifier、输出承诺", "证明私密输入满足转让约束并防重复消费"),
        ("外部事实层", "名册、银行、托管、KMS/HSM、OIDC、权威根发布", "由持牌/权威系统提供事实；平台验证、编排和留痕"),
    ], [1300, 3450, 4610], font_size=9.2)
    add_heading(doc, "3.2 两条业务轨道", 2)
    add_para(doc, "公开生命周期轨道用于认购、转让、赎回及外部名册/现金/托管对账；隐私结算轨道用于 2-in/2-out JoinSplit。两条轨道共享产品规则、身份、审计与外部事实边界，但不会把隐私 proof 直接冒充法定过户。")
    add_callout(doc, "隐私结算终局链",
                "服务端冻结 13 项公开输入 → Groth16 验证并原子占用 nullifier → ROOT_PENDING → operations maker 提交权威根与执行引用 → 不同 principal 的 checker 批准 → CONFIDENTIAL_NOTE_LEDGER / SETTLED → legalRegisterApplied=false，继续等待外部法定登记。",
                fill=BLUEGRAY, accent=NAVY)
    add_heading(doc, "3.3 为什么这套分层重要", 2)
    add_para(doc, "香港 SFC 2026 年关于代币化获认可投资产品的通函明确要求披露链上或链下哪一侧构成最终结算，并要求产品提供者对代币化安排和所有权记录承担最终责任。[1] BIS 也把结算终局定义为在法律上不可撤销、无条件且不因参与者破产而回转的时点。[2] 因此，平台把 proof、状态树、外部执行和法定名册拆开，不只是技术偏好，而是降低错误陈述和操作风险的必要控制。")

    add_heading(doc, "4. 主流产品比较", 1)
    add_heading(doc, "4.1 选择的对标对象", 2)
    add_table(doc, ["产品/网络", "公开定位与强项", "与我们的关系"], [
        ("Securitize", "发行、经纪、ATS、转让代理、基金行政等持牌全栈；已服务大型资管机构。[3]", "潜在合作方/强竞品；牌照、客户与分销远强于我们"),
        ("Tokeny / ERC-3643", "身份注册、资格声明、链上转让规则、发行与分销门户；标准化与互操作强。[4]", "链上合规与发行层强竞品；也可作为代币适配目标"),
        ("Taurus-CAPITAL", "多链、多标准、发行与资产服务，并与机构托管/HSM 集成。[5]", "发行、托管与多链覆盖强竞品；可作为托管/发行基础设施伙伴"),
        ("Canton Network", "面向机构的隐私与互操作网络，仅向相关参与者共享数据。[6]", "网络层与隐私理念相近；我们的优势在业务控制与终局编排"),
        ("本平台", "RWA 生命周期控制、原子账本、外部事实编排、角色化披露与 ZK JoinSplit", "聚焦跨系统可信流转；当前为 PoC，缺少牌照与生产生态"),
    ], [1550, 4800, 3010], font_size=8.7)

    add_heading(doc, "4.2 架构先进性：按维度给结论", 2)
    add_table(doc, ["维度", "公开市场领先者", "我们的相对判断"], [
        ("持牌全生命周期", "Securitize", "明显落后：我们是技术控制面，不持有同等牌照与渠道"),
        ("多链/多标准发行", "Taurus、Tokeny", "落后：当前没有同等级链与标准覆盖"),
        ("链上身份与合规标准", "Tokeny / ERC-3643", "互补偏后：我们规则更重链下机构流程，标准生态较弱"),
        ("机构隐私网络", "Canton", "各有侧重：Canton 强在网络互操作；我们强在 ZK 证明与业务终局编排"),
        ("账本与异步一致性", "未见统一公开基准", "设计较强：Serializable 事务、三账、Outbox、回调顺序、幂等与例外闭环"),
        ("终局分层与外部事实", "未见统一公开基准", "差异化强项：proof/根/执行/法定名册不混同，maker-checker 最终确认"),
        ("生产成熟度与客户证明", "头部现有平台", "明显落后：仍需正式审计、生产 KMS、客户连接、容量与灾备验证"),
    ], [2200, 2400, 4760], font_size=8.8)
    add_callout(doc, "最终判断：哪个架构更先进？",
                "没有单一冠军。若目标是立即发行、托管、分销和持牌运营，Securitize/Taurus/Tokeny 更成熟；若目标是跨机构隐私互操作，Canton 的网络生态更领先；若目标是把隐私证明、内部原子账本、外部执行与法律终局做严格隔离和可审计编排，我们的设计更聚焦、组合更完整。正确表述应是“在可信终局编排维度具有领先设计”，而不是“整体全球最先进”。",
                fill=PALE_GOLD, accent=GOLD)

    add_heading(doc, "5. 市场潜力评估", 1)
    add_heading(doc, "5.1 全球规模：高增长方向，但预测分歧巨大", 2)
    add_table(doc, ["来源/口径", "市场信号", "如何理解"], [
        ("RWA.xyz（2026-04-02 快照）", "分布式 RWA 约 276.5 亿美元；代表型资产约 4,413.8 亿美元。[7]", "真实链上分布规模仍处早期；不同口径不可直接相加"),
        ("McKinsey（2024 基准情景）", "2030 年代币化资产市值接近 2 万亿美元（不含加密资产）。[8]", "相对保守，强调基金、债券、贷款等先行"),
        ("BCG（2026 中性情景）", "2030 年约 14 万亿美元、2035 年约 55 万亿美元，且明确存在高度不确定性。[9]", "更宽口径、更激进；应作为上行情景而非预算依据"),
        ("DFSA 沙盒（2025）", "收到 96 份来自多地区的意向，覆盖债券、sukuk、基金与托管交易。[10]", "迪拜存在真实供给侧项目和监管承接能力"),
    ], [2600, 3150, 3610], font_size=8.7)
    add_para(doc, "预测跨度说明：2 万亿与 14 万亿美元并不互相证明或否定，差异来自资产范围、是否含现金/存款/稳定币、采用速度和“链上分布”与“数字表示”的定义。商业计划应以已签 PoC、上线产品数、受控资产规模和年度合同收入作为经营指标，而不是直接乘行业 TAM。", size=9.8, color=MUTED)

    add_heading(doc, "5.2 香港与迪拜：为什么是合理的首发区域", 2)
    add_table(doc, ["区域", "政策/市场信号", "产品对应机会"], [
        ("香港", "HKMA Ensemble Sandbox 自 2024 年测试代币化存款和资产；EnsembleTX 于 2025 年进入真实价值试点并在 2026 年持续。[11]", "基金、固定收益、流动性管理的机构控制面；强调所有权记录、终局披露和持牌分销"),
        ("迪拜/DIFC", "DFSA 已有 Investment Token 框架，并推进 Tokenisation Regulatory Sandbox；2026 规则继续强调授权、市场诚信与客户资产保护。[10][12]", "基金、债券、sukuk、私募资产的发行后控制、托管/名册接口和隐私机构结算"),
    ], [1200, 4300, 3860], font_size=9.0)
    add_heading(doc, "5.3 我们真正可服务的市场", 2)
    for x in [
        "第一优先：已有牌照、资产和客户，但缺少端到端技术编排的基金管理人、基金行政、分销平台和数字资产基础设施商。",
        "第二优先：希望在不公开交易金额和持仓关系的前提下，测试机构间转让与对账的银行、托管和资本市场基础设施。",
        "第三优先：拥有真实资产数据和法定名册，但现有系统无法提供可验证证据、跨系统幂等与异常闭环的资产发起方。",
        "不建议首发：面向无差别散户的开放交易市场。该路径同时要求牌照、分销、投资者保护、流动性和市场运营能力，超出当前产品边界。",
    ]:
        add_bullet(doc, x)
    add_callout(doc, "市场潜力评级：中高至高，但取决于机构接入",
                "技术可形成高价值 B2B 基础设施合同，且香港/迪拜的监管与试点方向吻合；但销售周期长、定制接口重、单个项目依赖多方。最有价值的护城河不是代码数量，而是逐步积累的持牌伙伴、权威数据连接、合规规则包、审计证据和 conformance test。")

    add_heading(doc, "6. 商业模式与客户切入", 1)
    add_heading(doc, "6.1 建议产品包装", 2)
    add_table(doc, ["产品层", "交付内容", "主要买方"], [
        ("RWA Control Plane", "角色、规则、状态机、账本、审计、异常与机构工作台", "资管、基金行政、发行平台"),
        ("Confidential Settlement", "JoinSplit 验证、nullifier、根管理、双人终局与最小披露", "银行、托管、机构交易网络"),
        ("Institution Connector Kit", "名册/银行/托管/NAV/身份/KMS 适配器与一致性测试", "系统集成商、持牌基础设施商"),
        ("Evidence & Oversight", "证据包、角色化监管视图、回执、对账与审计导出", "审计、合规、监管科技团队"),
    ], [1900, 4440, 3020], font_size=9.0)
    add_heading(doc, "6.2 收费结构", 2)
    for x in [
        "PoC/实施费：按单一产品、合成数据、明确接口和验收标准收费。",
        "平台订阅费：按机构、环境、产品数量、工作流模块和 SLA 收费。",
        "连接器与合规包：按银行/托管/名册/身份系统连接和司法辖区规则维护收费。",
        "交易或资产服务费：只在牌照、合同与业务模型允许时采用，不应把技术平台默认包装成经纪或交易场所。",
        "专业服务：部署、安全评估、灾备演练、证据导出与机构 conformance test。",
    ]:
        add_bullet(doc, x)
    add_heading(doc, "6.3 首个客户的最优 PoC", 2)
    add_number(doc, "选择一只香港或 DIFC 的合成基金/私募信贷/sukuk 场景，限定专业投资者与单一资产类型。")
    add_number(doc, "由合作方提供真实字段结构但使用脱敏或合成数据，明确法定名册、现金、托管和 NAV 的权威来源。")
    add_number(doc, "跑通认购、受控转让、赎回、凭证失效、名册超时、外部回调乱序、隐私转让与双人终局。")
    add_number(doc, "用共同签署的验收矩阵证明：失败不改余额、重复不重复入账、账本守恒、权限最小化、证据可导出。")
    add_number(doc, "PoC 结束后再讨论正式电路审计、生产连接、牌照边界和商业上线，不在演示阶段承诺真实资产终局。")

    add_heading(doc, "7. 当前成熟度、缺口与路线图", 1)
    add_heading(doc, "7.1 已有工程证据", 2)
    add_table(doc, ["项目", "当前状态"], [
        ("业务流程", "产品、角色、认购/转让/赎回、暂停、凭证、NAV、例外与 maker-checker 已实现"),
        ("数据与账本", "PostgreSQL 原子事务、资产账/名册账、一致性约束、审计链、Outbox、幂等已实现"),
        ("身份与数据保护", "OIDC/MFA/权限边界、CSRF、AES-GCM/KMS 接口、角色化脱敏视图已实现"),
        ("隐私结算", "真实本地 Groth16 JoinSplit 正反向量、13 项公开输入、nullifier、Merkle root 与两阶段终局已接线"),
        ("机构连接", "签名回调 HTTP 入口、Ed25519 Connector SDK、顺序/幂等/conformance 已实现"),
        ("隔离 Prover", "加密 witness 引用、持久化租约/重试、租户隔离、二次验证和运行指标已实现"),
        ("自动验收", "本地 PostgreSQL 最终验收 100/100 通过；生产依赖审计 0 个已知漏洞（截至本次验收）"),
        ("交付判定", "可交付合成数据机构 PoC；不可直接接真实客户资产"),
    ], [2300, 7060], font_size=9.2)
    add_heading(doc, "7.2 生产阻断项", 2)
    for x in [
        "正式 ZK 供应链：冻结电路、可复现构建、独立审计、正式 ceremony/setup、生产 vkey，并接入真实 Prover/Vault/HSM。",
        "权威根与执行引擎：机构提供可验证根发布批次及 fee/recipient/relayer 等外部执行回执。",
        "隐私存入/赎回自举：目前正式 ZK 轨道覆盖 2-in/2-out JoinSplit，初始 Note 铸造及退出需要目标结构。",
        "法定名册、现金与托管：必须接入持牌/权威事实源，并共同定义失败、回滚、替代与 SLA。",
        "生产 KMS/HSM、OIDC、网络、监控、灾备和独立渗透测试。",
        "香港/迪拜的产品结构、销售对象、数据责任、AML/KYC、托管及牌照法律意见。",
    ]:
        add_bullet(doc, x)
    add_heading(doc, "7.3 三阶段路线图", 2)
    add_table(doc, ["阶段", "目标", "退出标准"], [
        ("阶段 A：机构 PoC", "合成数据、单产品、单法域、接口模拟", "业务/安全/异常验收矩阵通过，合作方确认价值"),
        ("阶段 B：受控试点", "真实机构接口、受限参与者、独立审计", "生产身份/KMS、正式 ZK、灾备和法律边界签署"),
        ("阶段 C：生产扩展", "多产品、多机构、标准化连接器", "SLA、容量、监管报告、运营团队与持续审计成熟"),
    ], [1800, 3500, 4060], font_size=9.1)

    add_heading(doc, "8. 合作建议与最终结论", 1)
    add_heading(doc, "8.1 希望合作方提供", 2)
    for x in [
        "一个明确的资产场景、目标法域、投资者类别和持牌责任边界；",
        "真实但脱敏的数据字典、状态码、回调顺序、对账口径和异常案例；",
        "法定名册、银行、托管、NAV、身份和根发布的权威来源负责人；",
        "技术、安全、运营、合规和法律共同参与的 PoC 验收小组；",
        "对生产部署、数据驻留、密钥托管、SLA 与审计的约束。",
    ]:
        add_bullet(doc, x)
    add_heading(doc, "8.2 我们可以提供", 2)
    for x in [
        "可运行的 RWA 生命周期与隐私结算控制面；",
        "合成数据端到端演示、失败注入和 100 项自动验收基线；",
        "机构接口适配框架、签名回调、Outbox、幂等与异常处置工作流；",
        "以最小披露为原则的审计证据、监管视图和终局状态说明；",
        "从 PoC 到受控试点的安全、ZK、KMS、身份、连接器与部署路线图。",
    ]:
        add_bullet(doc, x)
    add_callout(doc, "最终结论",
                "这是一条值得继续投入的高潜力赛道，但成功公式不是“做出一个更复杂的代币”，而是“成为持牌机构之间可信状态流转的基础设施”。当前产品的技术内核已足以进入严肃 PoC；其最有竞争力的部分是终局分层、账本一致性、隐私转让和外部事实编排。下一阶段的价值增长将主要来自机构数据、权威连接、正式审计、法律结构与真实运营，而不是继续堆叠孤立功能。",
                fill=PALE_GOLD, accent=GOLD)

    doc.add_page_break()
    add_heading(doc, "附录 A：资料来源与说明", 1)
    sources = [
        ("[1] 香港证券及期货事务监察委员会，Circular on tokenisation of SFC-authorised investment products，2026-04-20。", "https://apps.sfc.hk/edistributionWeb/gateway/EN/circular/doc?refNo=26EC22"),
        ("[2] BIS/CPMI，Tokenisation in the context of money and other assets: concepts and implications for central banks，2024。", "https://www.bis.org/cpmi/publ/d225.htm"),
        ("[3] Securitize 官方材料：持牌全栈、基金行政与 tokenization 业务。", "https://investors.securitize.io/news/news-details/2025/Securitize-Acquires-MG-Stovers-Fund-Administration-Business-to-Become-the-Largest-Digital-Asset-Fund-Administrator-04-15-2025/default.aspx"),
        ("[4] Tokeny T-REX Platform / ERC-3643 官方文档。", "https://docs.tokeny.com/docs/t-rex-platform"),
        ("[5] Taurus-CAPITAL 官方产品页。", "https://www.taurushq.com/capital/"),
        ("[6] Canton Network 官方说明。", "https://www.canton.network/"),
        ("[7] RWA.xyz Global Market Overview，数据快照截至 2026-04-02。", "https://app.rwa.xyz/"),
        ("[8] McKinsey，The tides of tokenization / From ripples to waves，2024。", "https://www.mckinsey.com/featured-insights/charts/the-tides-of-tokenization"),
        ("[9] BCG，2026 Global Asset Management Report: An Imperative for Growth。", "https://www.bcg.com/publications/2026/an-imperative-for-growth-and-the-new-economics-of-asset-management"),
        ("[10] DFSA，Tokenisation Regulatory Sandbox 及 2025 年入选机构进展。", "https://www.dfsa.ae/news/dfsa-begins-engagement-firms-selected-its-tokenisation-regulatory-sandbox-reinforcing-its-commitment-responsible-innovation-difc"),
        ("[11] HKMA，Project Ensemble / EnsembleTX 2026 priorities。", "https://www.hkma.gov.hk/media/eng/publication-and-research/annual-report/2025/07_Priorities_for_2026_and_Beyond.pdf"),
        ("[12] DFSA，Crypto Token Regulation（规则于 2026-01-12 更新生效）。", "https://www.dfsa.ae/crypto"),
    ]
    for label, url in sources:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(label + " ")
        set_run_font(r, size=9.3, color=INK)
        add_hyperlink(p, "访问来源", url)
    add_para(doc, "说明：竞品信息来自各公司公开资料，未对其未披露的内部架构作负面推断；本平台状态来自本地代码、数据库迁移和最终验收报告。市场数字口径不同，不能直接相加。本文件不是法律、投资或监管意见，也不构成对收益、牌照或合规结果的保证。", size=9.2, color=MUTED, before=8, after=0)

    # Apply the same running furniture to all sections and set core properties.
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)
        add_header_footer(section)
    doc.core_properties.title = "RWA可信流转与隐私结算平台：产品介绍、竞品比较与市场潜力评估"
    doc.core_properties.subject = "合作方产品材料"
    doc.core_properties.author = "RWA Trust & Lifecycle Network"
    doc.core_properties.keywords = "RWA, Tokenization, Privacy Settlement, Zero Knowledge, Hong Kong, Dubai"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
